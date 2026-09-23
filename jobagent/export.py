from __future__ import annotations

from datetime import datetime
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from jobagent import db
from jobagent.display import telegram_post_url
from jobagent.paths import CSV_PATH, DATA_DIR, XLSX_PATH

HEADERS = [
    "Date",
    "Status",
    "Company",
    "Position",
    "Salary",
    "Remote",
    "Location",
    "Telegram",
    "Apply",
    "Post",
    "Channel",
    "How",
    "Notes",
]

COLUMN_WIDTHS = {
    "A": 20,
    "B": 12,
    "C": 28,
    "D": 38,
    "E": 42,
    "F": 10,
    "G": 16,
    "H": 18,
    "I": 42,
    "J": 36,
    "K": 22,
    "L": 12,
    "M": 22,
}

HEADER_FILL = PatternFill("solid", fgColor="1F4E5F")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
CELL_FONT = Font(name="Calibri", size=11, color="1F2A30")
LINK_FONT = Font(name="Calibri", size=11, color="0563C1", underline="single")
WRAP = Alignment(vertical="center", wrap_text=True, horizontal="left")
THIN = Border(
    left=Side(style="thin", color="D0D7DE"),
    right=Side(style="thin", color="D0D7DE"),
    top=Side(style="thin", color="D0D7DE"),
    bottom=Side(style="thin", color="D0D7DE"),
)
ZEBRA = PatternFill("solid", fgColor="F4F7F8")
STATUS_FILL = {
    "applied": PatternFill("solid", fgColor="C6EFCE"),
    "rejected": PatternFill("solid", fgColor="FFC7CE"),
    "interview": PatternFill("solid", fgColor="FFE699"),
}


def _raw(job: Any, key: str) -> Any:
    try:
        return job[key]
    except (IndexError, KeyError):
        return None


def _text(job: Any, key: str) -> str:
    try:
        value = job[key]
    except (IndexError, KeyError):
        value = None
    if value is None:
        return ""
    return str(value).strip()


def _pretty_date(iso: str) -> str:
    if not iso:
        return ""
    try:
        stamp = datetime.fromisoformat(iso.replace("Z", "+00:00"))
        return stamp.strftime("%d %b %Y  %H:%M UTC")
    except ValueError:
        return iso


def sheet_row(job: Any) -> list[str]:
    contact = _text(job, "telegram_contact").lstrip("@")
    post = telegram_post_url(
        _text(job, "channel_username") or None,
        _raw(job, "message_id"),
        _raw(job, "channel_id"),
    )
    return [
        _pretty_date(_text(job, "applied_at")),
        _text(job, "apply_status") or "applied",
        _text(job, "company"),
        _text(job, "position"),
        _text(job, "salary"),
        _text(job, "remote"),
        _text(job, "location"),
        f"@{contact}" if contact else "",
        _text(job, "apply_url"),
        post or "",
        _text(job, "channel_title"),
        _text(job, "apply_method"),
        _text(job, "notes"),
    ]


def _link_target(value: str, kind: str) -> str | None:
    if not value:
        return None
    if kind == "telegram":
        handle = value.lstrip("@")
        return f"https://t.me/{handle}" if handle else None
    if value.startswith("http://") or value.startswith("https://"):
        return value
    return None


def _style_workbook(book: Workbook, rows: list[list[str]]) -> None:
    sheet = book.active
    sheet.title = "Applied"
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{max(1, len(rows) + 1)}"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.row_dimensions[1].height = 22

    for index, header in enumerate(HEADERS, start=1):
        cell = sheet.cell(1, index, header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", horizontal="left")
        cell.border = THIN

    link_cols = {8: "telegram", 9: "url", 10: "url"}
    for row_index, values in enumerate(rows, start=2):
        sheet.row_dimensions[row_index].height = 36
        zebra = row_index % 2 == 0
        for col_index, value in enumerate(values, start=1):
            cell = sheet.cell(row_index, col_index, value)
            cell.font = CELL_FONT
            cell.alignment = WRAP
            cell.border = THIN
            if zebra:
                cell.fill = ZEBRA
            target = _link_target(value, link_cols.get(col_index, ""))
            if target:
                cell.hyperlink = target
                cell.font = LINK_FONT
        status = (values[1] or "applied").lower()
        status_cell = sheet.cell(row_index, 2)
        if status in STATUS_FILL:
            status_cell.fill = STATUS_FILL[status]

    for letter, width in COLUMN_WIDTHS.items():
        sheet.column_dimensions[letter].width = width


def export_sheets() -> tuple[str, str]:
    db.init_db()
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    with db.connect() as conn:
        jobs = db.export_rows(conn)
    rows = [sheet_row(job) for job in jobs]

    import csv

    with CSV_PATH.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(HEADERS)
        writer.writerows(rows)

    book = Workbook()
    _style_workbook(book, rows)
    try:
        book.save(XLSX_PATH)
    except PermissionError as exc:
        raise SystemExit(
            f"Close {XLSX_PATH} in LibreOffice or Excel, then try again."
        ) from exc
    return str(CSV_PATH), str(XLSX_PATH)
